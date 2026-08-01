"""Entrypoint that runs the model-authored build script and reports what it produced.

The worker parses the final stdout line, so the manifest marker must be the last thing written
to stdout. Everything the build script prints stays above it and is surfaced to the user as
build logs.

Exit codes are distinct because the Go orchestrator treats them differently: a script error is
worth a self-heal retry, a verification failure usually is too, but an empty /out means the
script ran to completion while misunderstanding the contract, which a retry rarely fixes.
"""

import json
import os
import runpy
import sys
import traceback

WORKSPACE = "/workspace"
OUT_DIR = os.path.join(WORKSPACE, "out")
SCRIPT = os.path.join(WORKSPACE, "build.py")
MANIFEST_MARKER = "__BOXAI_MANIFEST__"

EXIT_SCRIPT_ERROR = 1
EXIT_NO_ARTIFACTS = 2
EXIT_VERIFICATION_FAILED = 3


def verify_docx(path):
    import docx

    document = docx.Document(path)
    # Truth-testing an lxml element is deprecated and counts children, which is not the question
    # being asked here.
    if document.element.body is None:
        raise ValueError("document body is empty")


def verify_xlsx(path):
    import openpyxl

    workbook = openpyxl.load_workbook(path)
    if not workbook.sheetnames:
        raise ValueError("workbook has no sheets")


def verify_pptx(path):
    import pptx

    presentation = pptx.Presentation(path)
    if len(presentation.slides) == 0:
        raise ValueError("presentation has no slides")


def verify_pdf(path):
    import pypdf

    reader = pypdf.PdfReader(path)
    if len(reader.pages) == 0:
        raise ValueError("pdf has no pages")


VERIFIERS = {
    ".docx": verify_docx,
    ".xlsx": verify_xlsx,
    ".pptx": verify_pptx,
    ".pdf": verify_pdf,
}


def collect_artifacts():
    """Reopen every produced file with the library that owns its format.

    A script can write a file that is structurally corrupt and still exit zero — python-docx in
    particular will happily serialize a tree the Word renderer rejects. Reopening is the cheapest
    check that catches it before the user downloads something that will not open.
    """
    artifacts = []
    failures = []
    for root, _, names in os.walk(OUT_DIR):
        for name in sorted(names):
            if name.startswith("."):
                continue
            absolute = os.path.join(root, name)
            relative = os.path.relpath(absolute, OUT_DIR)
            extension = os.path.splitext(name)[1].lower()
            entry = {
                "name": relative,
                "bytes": os.path.getsize(absolute),
                "verified": False,
            }
            verifier = VERIFIERS.get(extension)
            if verifier is None:
                artifacts.append(entry)
                continue
            try:
                verifier(absolute)
                entry["verified"] = True
            except Exception as err:
                failures.append(f"{relative}: {type(err).__name__}: {err}")
            artifacts.append(entry)
    return artifacts, failures


def emit(payload, code):
    sys.stdout.flush()
    print(f"{MANIFEST_MARKER} {json.dumps(payload, ensure_ascii=False)}")
    sys.stdout.flush()
    sys.exit(code)


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    if not os.path.exists(SCRIPT):
        print(f"{SCRIPT} was never written", file=sys.stderr)
        emit({"artifacts": [], "error": "missing build script"}, EXIT_SCRIPT_ERROR)

    os.chdir(WORKSPACE)
    try:
        runpy.run_path(SCRIPT, run_name="__main__")
    except SystemExit as err:
        if err.code not in (0, None):
            traceback.print_exc()
            emit({"artifacts": [], "error": f"script exited with {err.code}"}, EXIT_SCRIPT_ERROR)
    except BaseException:
        traceback.print_exc()
        emit({"artifacts": [], "error": "script raised"}, EXIT_SCRIPT_ERROR)

    artifacts, failures = collect_artifacts()
    if not artifacts:
        print(f"build script wrote nothing to {OUT_DIR}", file=sys.stderr)
        emit({"artifacts": [], "error": "no artifacts"}, EXIT_NO_ARTIFACTS)
    if failures:
        for failure in failures:
            print(f"verification failed: {failure}", file=sys.stderr)
        emit({"artifacts": artifacts, "error": "; ".join(failures)}, EXIT_VERIFICATION_FAILED)

    emit({"artifacts": artifacts, "error": ""}, 0)


if __name__ == "__main__":
    main()
